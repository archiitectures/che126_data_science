"""FSRI instrument data and permissioned structured-file imports."""

import csv
import io
import json
import re
from pathlib import Path

from che126_data_science.data_science.master import make_observation

from .fsri import parse_hfm_results
from .normalization import canonical_unit_text


def decode_instrument(data: bytes) -> str:
    return data.decode(
        "utf-16" if data[:2] in {b"\xff\xfe", b"\xfe\xff"} else "utf-8-sig"
    )


def _hfm(document, text, filename, *, locator=None, context=None):
    context = context or {}
    kind = (
        "conductivity"
        if "conductivity" in filename.lower()
        else "heatcapacity"
        if "heatcapacity" in filename.lower()
        else None
    )
    if not kind:
        raise ValueError("Unknown HFM instrument measurement")
    marker = "Results Table -- SI Units"
    if marker not in text:
        # Pinned excerpt format is still supported, but complete .tst requires marker.
        if not context.get("is_excerpt"):
            raise ValueError("HFM instrument file lacks SI results marker")
        table = text
    else:
        table = text.split(marker, 1)[1].split("Results Table --", 1)[0]
    observations = parse_hfm_results(table, kind)
    material = context.get("polymer") or filename.split("_", 1)[0]
    match = re.search(r"(?:^|_)(R\d+)(?:\.|$)", filename, re.IGNORECASE)
    replicate = match[1] if match else None
    state = (
        "Wet"
        if re.search(r"_Wet_", filename, re.IGNORECASE)
        else "Dry"
        if re.search(r"_Dry_", filename, re.IGNORECASE)
        else "not_reported"
    )
    prop, unit = (
        ("thermal_conductivity", "W/(m K)")
        if kind == "conductivity"
        else ("volumetric_heat_capacity", "J/(m3 K)")
    )
    rows = [
        make_observation(
            document,
            label=kind,
            prop=prop,
            raw_value=p.value,
            raw_unit=unit,
            sample_id=filename,
            source_table="Results Table -- SI Units",
            source_row=f"{locator or filename}:row {p.table_row}",
            polymer=material,
            sample_form="bulk_reference",
            material_class="bulk polymer reference",
            temperature_C=p.temperature_C,
            temperature_role="measurement",
            replicate_id=replicate,
            direction="through-plane",
            method="TA Instruments FOX HFM; SI results",
            measurement_context=f"FSRI {state}; ambient/unconditioned Wet is not electrolyte soaking",
            record_type="measured",
            evidence_group=f"fsri:{filename}:si:{p.table_row}",
            extraction_method="fsri_hfm_si",
            notes="Instrument-reported results; repeated setpoints retained as rows, not independent experiments.",
        )
        for p in observations
    ]
    return rows, {
        "kind": "instrument_table",
        "filename": filename,
        "locator": locator,
        "text": text,
        "results_si": table,
    }


def extract_fsri(document, context):
    data = Path(document.path).read_bytes()
    filename = context.get("filename") or Path(document.url or document.path).name
    if filename.endswith(".json"):
        snapshot = json.loads(data)
        if "exports" not in snapshot:
            return [], [{"kind": "metadata", "data": snapshot}]
        if not re.fullmatch(r"[a-f0-9]{40}", snapshot.get("source_commit", "")):
            raise ValueError("FSRI snapshot requires a pinned commit")
        rows, artifacts = [], []
        for entry in snapshot["exports"]:
            path = entry["path"]
            if "results_si" in entry:
                new, artifact = _hfm(
                    document,
                    entry["results_si"],
                    Path(path).name,
                    locator=f"{snapshot['source_commit']}/{path}; blob {entry['blob_sha']}",
                    context={**context, "is_excerpt": True},
                )
                rows.extend(new)
                artifacts.append(artifact)
            elif "csv" in entry:
                new, artifact = _density(
                    document, entry["csv"], Path(path).name, locator=path
                )
                rows.extend(new)
                artifacts.append(artifact)
        return rows, artifacts
    text = decode_instrument(data)
    if filename.lower().endswith(".tst"):
        rows, artifact = _hfm(document, text, filename, context=context)
        return rows, [artifact]
    if "density" in filename.lower():
        rows, artifact = _density(document, text, filename)
        return rows, [artifact]
    return extract_manual(document, context)


def _density(document, text, filename, locator=None):
    cells = list(csv.reader(io.StringIO(text)))
    matches = [
        (i, row)
        for i, row in enumerate(cells, 1)
        if row and row[0].strip() == "Sample Density [kg/m3]"
    ]
    if len(matches) != 1 or len(matches[0][1]) != 2:
        raise ValueError("Unsupported density CSV")
    i, row = matches[0]
    return [
        make_observation(
            document,
            label=row[0],
            prop="density",
            raw_value=row[1],
            raw_unit="kg/m3",
            sample_id=filename,
            source_row=f"{locator or filename}:row {i}",
            polymer=filename.split("_")[0],
            sample_form="bulk_reference",
            material_class="bulk polymer reference",
            record_type="measured",
            replicate_id=Path(filename).stem.rsplit("_", 1)[-1],
            evidence_group=f"fsri:{filename}:density",
            extraction_method="fsri_density_csv",
            notes="Reference density has no reported measurement temperature. Identity from source directory/filename, not possibly mislabeled CSV column.",
        )
    ], {"kind": "csv", "filename": filename, "cells": cells}


def extract_manual(document, context):
    """Canonical or explicitly mapped CSV/TSV/XLSX, without guessing field names."""
    filename = context.get("filename") or document.path
    if str(filename).lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        workbook = load_workbook(document.path, read_only=True, data_only=True)
        try:
            sheet = workbook[context.get("sheet", workbook.sheetnames[0])]
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, None)
            if not header:
                raise ValueError("Workbook sheet has no header")
            source_rows = [dict(zip(header, row, strict=True)) for row in iterator]
        finally:
            workbook.close()
    else:
        text = Path(document.path).read_text(encoding="utf-8-sig")
        source_rows = list(
            csv.DictReader(
                io.StringIO(text),
                delimiter="\t" if str(filename).lower().endswith(".tsv") else ",",
            )
        )
    mapping = context.get("columns", {})
    rows = []
    for i, raw in enumerate(source_rows, 2):
        try:
            source = {mapping.get(k, k): v for k, v in raw.items()}
            source = {
                k: v
                for k, v in source.items()
                if v is not None and str(v).strip() != ""
            }
            required = {"property", "value", "unit", "sample_id"}
            if not required.issubset(source):
                raise ValueError(
                    f"Structured row {i} missing {sorted(required - source.keys())}; supply explicit column mapping"
                )
            meta = {
                k: source[k]
                for k in (
                    "polymer",
                    "manufacturer",
                    "product",
                    "separator_architecture",
                    "coating",
                    "sample_form",
                    "electrolyte_salt",
                    "electrolyte_solvent",
                    "electrolyte_concentration",
                    "direction",
                    "method",
                    "measurement_context",
                    "replicate_id",
                    "curve_id",
                    "shrinkage_basis",
                    "source_url",
                    "source_table",
                    "source_column",
                    "notes",
                )
                if k in source
            }
            for k in (
                "temperature_C",
                "duration_min",
                "heating_rate_C_min",
                "uncertainty",
            ):
                if k in source:
                    meta[k] = float(source[k])
            legacy = "record_id" in source or "value_type" in source
            if legacy:
                meta["quality_flag"] = "legacy_primary_source_verification_required"
                meta["polymer"] = source.get("material")
                meta["measurement_context"] = source.get("conditions")
                meta["source_table"] = source.get("source_locator")
            state = source.get("electrolyte_state", source.get("state", "not_reported"))
            record_type = source.get("record_type", "unknown")
            rows.append(
                make_observation(
                    document,
                    label=str(
                        source.get("original_property_label", source["property"])
                    ),
                    prop=str(source["property"]),
                    raw_value=source["value"],
                    raw_unit=str(source["unit"]),
                    normalization_unit=canonical_unit_text(str(source["unit"])),
                    sample_id=str(source["sample_id"]),
                    source_row=str(i),
                    statistic=str(source.get("statistic", "reported")),
                    record_type=str(record_type),
                    temperature_role=source.get("temperature_role", "unknown"),
                    electrolyte_state="not_reported" if state == "unknown" else state,
                    extraction_method="manual_structured_import",
                    **meta,
                )
            )
        except (ValueError, TypeError, KeyError) as exc:
            rows.append(
                make_observation(
                    document,
                    label=str(raw.get("property") or "Invalid structured row"),
                    prop="unresolved",
                    raw_value=json.dumps(raw, default=str),
                    raw_unit="",
                    sample_id=str(raw.get("sample_id") or f"row_{i}"),
                    source_row=str(i),
                    quality_flag="invalid_structured_row",
                    notes=str(exc),
                    extraction_method="manual_structured_import",
                )
            )
    return rows, [{"kind": "structured_rows", "rows": source_rows}]
